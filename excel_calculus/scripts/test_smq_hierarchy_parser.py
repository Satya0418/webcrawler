import openpyxl

path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/MedDRA SMQ list/SMQ_spreadsheet_29_0_English.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["SMQ 29.0"]

# Robust hierarchy parser
smq_terms = {} # smq_name -> {"Narrow": set(), "Broad": set()}
stack = [] # list of (col_idx, smq_name)

for r in range(2, ws.max_row + 1):
    row_vals = [ws.cell(r, c).value for c in range(1, 15)]
    
    # Check if this row defines an SMQ
    smq_col = -1
    smq_name = None
    for c_idx, val in enumerate(row_vals):
        if val and "(smq)" in str(val).lower():
            smq_col = c_idx
            smq_name = str(val).strip()
            break
            
    if smq_name:
        # Pop stack items at or deeper than this column
        while stack and stack[-1][0] >= smq_col:
            stack.pop()
        stack.append((smq_col, smq_name))
        if smq_name not in smq_terms:
            smq_terms[smq_name] = {"Narrow": set(), "Broad": set()}
        continue
        
    # Check if this row is a PT row
    # Look for 'Narrow' or 'Broad'
    scope_col = -1
    scope_val = None
    for c_idx, val in enumerate(row_vals):
        if val in ("Narrow", "Broad"):
            scope_col = c_idx
            scope_val = val
            break
            
    if scope_val and scope_col >= 0:
        # Next col is status, then PT name, then PT code
        # Wait, sometimes status is skipped or at c_idx+1
        # Let us check c_idx+1, c_idx+2
        status = row_vals[scope_col + 1]
        pt_name = row_vals[scope_col + 2]
        pt_code = row_vals[scope_col + 3]
        
        if pt_name:
            pt_clean = str(pt_name).strip()
            # Add to all SMQs currently in the stack
            for _, s_name in stack:
                if s_name not in smq_terms:
                    smq_terms[s_name] = {"Narrow": set(), "Broad": set()}
                smq_terms[s_name]["Broad"].add(pt_clean)
                if scope_val == "Narrow":
                    smq_terms[s_name]["Narrow"].add(pt_clean)

print(f"Total parsed SMQs: {len(smq_terms)}")

# Check our target SMQs
target_smqs = [
    "Drug related hepatic disorders - comprehensive search (SMQ)",
    "Osteoporosis/osteopenia (SMQ)",
    "Rhabdomyolysis/myopathy (SMQ)",
    "Medication errors (SMQ)"
]

for t in target_smqs:
    data = smq_terms.get(t, {"Narrow": set(), "Broad": set()})
    print(f"\nTarget: {t}")
    print(f"  Narrow terms count: {len(data['Narrow'])}")
    print(f"  Broad terms count: {len(data['Broad'])}")
    sample = list(data['Broad'])[:5]
    print(f"  Sample terms: {sample}")
