import openpyxl
import re

# Load SMQ workbook
smq_path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/MedDRA SMQ list/SMQ_spreadsheet_29_0_English.xlsx"
wb_smq = openpyxl.load_workbook(smq_path, data_only=True)
ws_smq = wb_smq["SMQ 29.0"]

# Parse SMQs
smq_dict = {} # (smq_name, scope) -> set of lowercase PT names
current_smq = None
for r in range(2, ws_smq.max_row + 1):
    val0 = ws_smq.cell(r, 1).value
    val1 = ws_smq.cell(r, 2).value # Scope (Narrow / Broad)
    val3 = ws_smq.cell(r, 4).value # PT Name
    if val0 and "(SMQ)" in str(val0):
        current_smq = str(val0).strip()
    elif current_smq and val1 in ("Narrow", "Broad") and val3:
        pt = str(val3).strip().lower()
        if (current_smq, "Broad") not in smq_dict:
            smq_dict[(current_smq, "Broad")] = set()
        smq_dict[(current_smq, "Broad")].add(pt)
        
        if val1 == "Narrow":
            if (current_smq, "Narrow") not in smq_dict:
                smq_dict[(current_smq, "Narrow")] = set()
            smq_dict[(current_smq, "Narrow")].add(pt)

print(f"Parsed {len(smq_dict)} SMQ scope entries.")

# Check hepatic comprehensive search
hepatic_broad_keys = [k for k in smq_dict.keys() if "hepatic disorders - comprehensive" in k[0].lower()]
print("Hepatic keys found:", hepatic_broad_keys)
hepatic_broad_pts = smq_dict.get(hepatic_broad_keys[0]) if hepatic_broad_keys else set()
print("Hepatic broad PT count:", len(hepatic_broad_pts))

# Check osteoporosis
osteo_keys = [k for k in smq_dict.keys() if "osteoporosis" in k[0].lower()]
print("Osteo keys found:", osteo_keys)

# Check rhabdomyolysis
rhabdo_keys = [k for k in smq_dict.keys() if "rhabdomyolysis" in k[0].lower()]
print("Rhabdo keys found:", rhabdo_keys)

# Check medication error
med_error_keys = [k for k in smq_dict.keys() if "medication error" in k[0].lower()]
print("Medication error keys found:", med_error_keys)

# Now load Abiraterone Interval LL
ll_path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Abiraterone/Abiraterone_20260428_CAN-KUW-OMAN-UAE PBRER_Interval Linelisting.xlsx"
wb_ll = openpyxl.load_workbook(ll_path, data_only=True)
ws_ll = wb_ll.active

rows = list(ws_ll.iter_rows(values_only=True))
headers = rows[0]
data = rows[1:]

idx_case = headers.index("Case Number")
idx_ev = headers.index("Event Verbatim")
idx_soc = headers.index("System Organ Class")
idx_outcome = headers.index("Outcome of Event")
idx_prod = headers.index("Product Name")

# Normalize event terms from Event Verbatim
case_events = []
bracket_pattern = re.compile(r'\[([^\]]+)\]')

for r_idx, r in enumerate(data):
    c_num = r[idx_case]
    ev_raw = str(r[idx_ev]) if r[idx_ev] else ""
    ev_clean = re.sub(r'_x[0-9a-fA-F]{4}_', '', ev_raw)
    bracket_terms = bracket_pattern.findall(ev_clean)
    
    soc = str(r[idx_soc]) if r[idx_soc] else ""
    prod = str(r[idx_prod]) if r[idx_prod] else ""
    
    for pos, term in enumerate(bracket_terms):
        term_clean = re.sub(r'[\r\n\t]+', ' ', term)
        term_clean = re.sub(r'\s+', ' ', term_clean).strip()
        if term_clean:
            case_events.append({
                "case_number": c_num,
                "raw_event": term,
                "clean_event": term_clean,
                "clean_event_lower": term_clean.lower(),
                "soc": soc,
                "product": prod,
                "row_idx": r_idx + 2,
                "pos": pos
            })

print(f"\nTotal extracted case-events: {len(case_events)}")

# TEST 1: Hepatotoxicity (Broad SMQ)
hep_matches = [e for e in case_events if e["clean_event_lower"] in hepatic_broad_pts]
hep_cases = sorted(set(e["case_number"] for e in hep_matches))
print(f"\nHepatotoxicity Matches: {len(hep_matches)} events across {len(hep_cases)} distinct cases")
for c in hep_cases[:5]:
    matched_evs = [e["clean_event"] for e in hep_matches if e["case_number"] == c]
    print(f"  Case {c}: {matched_evs}")

# TEST 2: Cardiac Disorders (SOC)
cardiac_cases = sorted(set(r[idx_case] for r in data if r[idx_soc] and "cardiac disorder" in str(r[idx_soc]).lower()))
print(f"\nCardiac Disorders (SOC match): {len(cardiac_cases)} distinct cases")
print("  Cases:", cardiac_cases)

# TEST 3: Rhabdomyolysis / Myopathy (Narrow SMQ)
rhabdo_narrow_pts = smq_dict.get(('Rhabdomyolysis/myopathy (SMQ)', 'Narrow'), set())
rhabdo_matches = [e for e in case_events if e["clean_event_lower"] in rhabdo_narrow_pts]
rhabdo_cases = sorted(set(e["case_number"] for e in rhabdo_matches))
print(f"\nRhabdomyolysis / Myopathy Matches: {len(rhabdo_matches)} events across {len(rhabdo_cases)} distinct cases")
print("  Cases:", rhabdo_cases)

# TEST 4: Cataract (PT list)
cataract_pts = {"cataract", "atopic cataract", "cataract nuclear", "cataract cortical", "toxic cataract", "cataract subcapsular"}
cat_matches = [e for e in case_events if e["clean_event_lower"] in cataract_pts]
cat_cases = sorted(set(e["case_number"] for e in cat_matches))
print(f"\nCataract Matches: {len(cat_matches)} events across {len(cat_cases)} distinct cases")
print("  Cases:", cat_cases)

# TEST 5: CYP2D6 (PT list)
cyp_pts = {"drug interaction", "potentiating drug interaction", "labelled drug-drug interaction issue", "labelled drug-drug interaction medication error"}
cyp_matches = [e for e in case_events if e["clean_event_lower"] in cyp_pts]
cyp_cases = sorted(set(e["case_number"] for e in cyp_matches))
print(f"\nCYP2D6 initial PT Matches: {len(cyp_matches)} events across {len(cyp_cases)} distinct cases")
for c in cyp_cases:
    matched_evs = [e["clean_event"] for e in cyp_matches if e["case_number"] == c]
    prod = [e["product"] for e in cyp_matches if e["case_number"] == c][0]
    print(f"  Case {c}: Events: {matched_evs} | Product: {prod}")
