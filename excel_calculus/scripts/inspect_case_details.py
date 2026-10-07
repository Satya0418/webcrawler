import openpyxl
import json

path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Abiraterone/Abiraterone_20260428_CAN-KUW-OMAN-UAE PBRER_Interval Linelisting.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb.active

headers = [cell for cell in next(ws.iter_rows(values_only=True))]
print("HEADERS count:", len(headers))
for i, h in enumerate(headers):
    print(f"Col {i}: {h}")

print("\n" + "="*50)
for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
    if r_idx == 0:
        continue
    if r_idx > 5:
        break
    case_num = row[0]
    ev = row[12] # Event Verbatim
    prod = row[21] # Product Name
    outcome_ev = row[23] # Outcome of Event
    soc = row[1]
    print(f"\n--- CASE: {case_num} ---")
    print(f"SOC: {soc}")
    print(f"EVENT VERBATIM RAW:\n{repr(ev)}")
    print(f"EVENT VERBATIM DISPLAY:\n{ev}")
    print(f"OUTCOME OF EVENT:\n{repr(outcome_ev)}")
    print(f"PRODUCT NAME:\n{repr(prod)}")
