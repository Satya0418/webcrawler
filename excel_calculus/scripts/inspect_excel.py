import openpyxl

files = {
    "SMQ": "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/MedDRA SMQ list/SMQ_spreadsheet_29_0_English.xlsx",
    "Abiraterone_LL": "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Abiraterone/Abiraterone_20260428_CAN-KUW-OMAN-UAE PBRER_Interval Linelisting.xlsx",
    "Oxycodone_LL": "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Oxycodone/Oxycodone_20260412_CAN PBRER_Interval LL.xlsx"
}

for name, path in files.items():
    print(f"\n{'='*20} {name} {'='*20}")
    wb = openpyxl.load_workbook(path, read_only=True)
    print("Sheets:", wb.sheetnames)
    for sheet_name in wb.sheetnames[:3]:
        ws = wb[sheet_name]
        print(f"\n--- Sheet: {sheet_name} ---")
        rows = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx < 10:
                rows.append([str(c) if c is not None else "" for c in row])
            else:
                break
        print(f"Total rows inspected: {len(rows)}")
        for i, r in enumerate(rows[:5]):
            # truncate long cell values for summary
            summary = [c[:30] + "..." if len(c) > 30 else c for c in r]
            print(f"Row {i}: {summary}")
