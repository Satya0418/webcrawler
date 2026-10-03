import docx
import json

doc_path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Abiraterone/Abiraterone CAN PBRER_DLP_20260428_Safety concerns and search criteria_v0.4_Final.docx"
doc = docx.Document(doc_path)

print("PARAGRAPHS:")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")

print("\n" + "="*50 + "\nTABLES:")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- TABLE {t_idx} (rows: {len(table.rows)}, cols: {len(table.columns)}) ---")
    for r_idx, row in enumerate(table.rows):
        row_vals = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
        print(f"R{r_idx}: {row_vals}")
