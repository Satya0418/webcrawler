import openpyxl

abi_path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Abiraterone/Abiraterone_20260428_CAN-KUW-OMAN-UAE PBRER_Interval Linelisting.xlsx"
oxy_path = "/Users/satya/Downloads/Required Artifacts for section 16.3 (3)/Oxycodone/Oxycodone_20260412_CAN PBRER_Interval LL.xlsx"

wb_abi = openpyxl.load_workbook(abi_path, data_only=True)
ws_abi = wb_abi.active
abi_cols = [c for c in next(ws_abi.iter_rows(values_only=True)) if c is not None]

wb_oxy = openpyxl.load_workbook(oxy_path, data_only=True)
ws_oxy = wb_oxy.active
oxy_cols = [c for c in next(ws_oxy.iter_rows(values_only=True)) if c is not None and str(c).strip()]

print(f"Abiraterone columns ({len(abi_cols)}):")
for i, c in enumerate(abi_cols):
    print(f"  {i+1}. {c}")

print(f"\nOxycodone columns ({len(oxy_cols)}):")
for i, c in enumerate(oxy_cols):
    print(f"  {i+1}. {c}")

common = set(abi_cols).intersection(set(oxy_cols))
print(f"\nCommon columns ({len(common)}):")
for c in sorted(common):
    print(f"  * {c}")

abi_only = set(abi_cols) - set(oxy_cols)
print(f"\nAbiraterone only ({len(abi_only)}):")
for c in sorted(abi_only):
    print(f"  - {c}")

oxy_only = set(oxy_cols) - set(abi_cols)
print(f"\nOxycodone only ({len(oxy_only)}):")
for c in sorted(oxy_only):
    print(f"  + {c}")
