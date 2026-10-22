import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelReportService:
    def __init__(self, db_session=None):
        self.db = db_session

    def generate_section_16_1_excel(self, report_data: dict) -> bytes:
        """
        Generates a submission-ready, formally styled Microsoft Excel (.xlsx) workbook
        representing the PBRER Section 16.1 Summary of Safety Concerns table.
        Strictly enforces:
        - 2-column layout (Risk Term, Number of Relevant Case Reports)
        - Clean category grouping rows
        - No total row across safety concerns
        - Professional typography, borders, and column sizing
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Section 16.1 Safety Summary"

        # Ensure grid lines are visible
        ws.views.sheetView[0].showGridLines = True

        product_name = report_data.get("product_name", "Unknown Product")
        reporting_period = report_data.get("reporting_period", "")

        # Fonts
        font_title = Font(name="Calibri", size=13, bold=True, color="0F172A")
        font_subtitle = Font(name="Calibri", size=10, italic=True, color="475569")
        font_th = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        font_category = Font(name="Calibri", size=10, bold=True, color="0F172A")
        font_risk = Font(name="Calibri", size=10, color="1E293B")
        font_count = Font(name="Calibri", size=10, bold=True, color="0F172A")

        # Fills
        fill_th = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        fill_category = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        # Alignments
        align_th_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_th_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_cat = Alignment(horizontal="left", vertical="center")
        align_risk = Alignment(horizontal="left", vertical="center", indent=1)
        align_count = Alignment(horizontal="center", vertical="center")

        # Borders
        thin_side = Side(border_style="thin", color="CBD5E1")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        border_th = Border(
            left=Side(border_style="thin", color="475569"),
            right=Side(border_style="thin", color="475569"),
            top=Side(border_style="medium", color="0F172A"),
            bottom=Side(border_style="medium", color="0F172A")
        )

        # Title Block
        ws["A1"] = "PBRER Section 16.1 — Summary of Safety Concerns"
        ws["A1"].font = font_title
        ws.row_dimensions[1].height = 24

        ws["A2"] = f"Product: {product_name} | Reporting Period: {reporting_period}"
        ws["A2"].font = font_subtitle
        ws.row_dimensions[2].height = 18

        # Blank Row 3
        ws.row_dimensions[3].height = 10

        # Table Header Row 4
        ws["A4"] = "Risk Term"
        ws["B4"] = "Number of Relevant Case Reports"

        for col, align in (("A4", align_th_left), ("B4", align_th_center)):
            ws[col].font = font_th
            ws[col].fill = fill_th
            ws[col].alignment = align
            ws[col].border = border_th
        ws.row_dimensions[4].height = 26

        curr_row = 5

        for section in report_data.get("table_sections", []):
            cat_name = section.get("category_name", "").upper()

            # Category Banner Row spanning A & B
            ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
            cell_cat = ws.cell(row=curr_row, column=1, value=cat_name)
            cell_cat.font = font_category
            cell_cat.fill = fill_category
            cell_cat.alignment = align_cat

            cell_cat_b = ws.cell(row=curr_row, column=2)
            cell_cat_b.fill = fill_category

            for col in range(1, 3):
                ws.cell(row=curr_row, column=col).border = border_all
            ws.row_dimensions[curr_row].height = 20
            curr_row += 1

            for risk in section.get("risks", []):
                risk_term = risk.get("risk_term", "")
                rel_count = risk.get("number_of_relevant_cases", 0)

                cell_r = ws.cell(row=curr_row, column=1, value=risk_term)
                cell_r.font = font_risk
                cell_r.alignment = align_risk
                cell_r.border = border_all
                cell_r.fill = fill_white

                cell_c = ws.cell(row=curr_row, column=2, value=int(rel_count))
                cell_c.font = font_count
                cell_c.alignment = align_count
                cell_c.border = border_all
                cell_c.fill = fill_white
                cell_c.number_format = "#,##0"

                ws.row_dimensions[curr_row].height = 20
                curr_row += 1

        # Adjust column widths
        ws.column_dimensions["A"].width = 54
        ws.column_dimensions["B"].width = 30

        output = io.BytesIO()
        wb.save(output)
        wb.close()
        return output.getvalue()
