import os
import openpyxl
from sqlalchemy.orm import Session
from excel_calculus.backend.app.models.entities import SMQTerm, AuditLog, Dataset
import hashlib
from datetime import datetime

class SMQLoaderService:
    def __init__(self, db: Session):
        self.db = db

    def load_smq_file(self, file_path: str, reference_version: str = "29.0"):
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"SMQ reference file not found: {file_path}")

        file_name = os.path.basename(file_path)
        
        # Calculate file hash for lineage
        with open(file_path, "rb") as f:
            file_hash = hashlib.sha256(f.read()).hexdigest()

        # Load workbook in read_only data_only mode for maximum speed
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        sheet_names = wb.sheetnames

        # Step 1: Map PT Code -> PT Name from the primary SMQ definition sheet
        # Usually 'SMQ 29.0' or sheet with 'smq' in name
        smq_pt_sheet = None
        for s in sheet_names:
            s_lower = s.lower()
            if "smq" in s_lower and "hierarchy" not in s_lower and "summary" not in s_lower and "count" not in s_lower and "code" not in s_lower:
                smq_pt_sheet = s
                break
        if not smq_pt_sheet:
            smq_pt_sheet = "SMQ 29.0" if "SMQ 29.0" in sheet_names else sheet_names[0]

        ws_smq = wb[smq_pt_sheet]
        code_to_name = {}
        code_to_category = {}

        for row in ws_smq.iter_rows(values_only=True):
            if not row:
                continue
            for c in range(len(row) - 3):
                val = row[c]
                if val in ("Narrow", "Broad"):
                    category = row[c + 1] if c + 1 < len(row) else None
                    pt_name = row[c + 2] if c + 2 < len(row) else None
                    pt_code = row[c + 3] if c + 3 < len(row) else None
                    if pt_code and pt_name:
                        c_str = str(pt_code).strip()
                        n_str = str(pt_name).strip()
                        code_to_name[c_str] = n_str
                        if category:
                            code_to_category[c_str] = str(category).strip()

        # Step 2: Load Active PT Codes (the authoritative MedDRA MSSO pre-resolved hierarchy sheet)
        has_active_codes_sheet = "Active PT Codes" in sheet_names

        batch = []
        batch_size = 5000
        total_inserted = 0
        seen_keys = set()  # (smq_name, scope, pt_code) to ensure uniqueness

        # Clear existing SMQ terms for this reference version
        self.db.query(SMQTerm).filter_by(reference_version=reference_version).delete()

        if has_active_codes_sheet:
            ws_codes = wb["Active PT Codes"]
            for row in ws_codes.iter_rows(values_only=True):
                if not row or len(row) < 4:
                    continue
                code_raw = row[0]
                smq_name_raw = row[1]
                scope_raw = row[2]
                pt_codes_raw = row[3]

                if not code_raw or not str(code_raw).strip().isdigit() or not smq_name_raw or not pt_codes_raw:
                    continue

                smq_code = str(code_raw).strip()
                smq_name = str(smq_name_raw).strip()
                raw_scope = str(scope_raw).strip()

                # MedDRA defines 'Narrow' for specific search, 'Broad+Narrow' for comprehensive broad search
                scope = "Broad" if raw_scope == "Broad+Narrow" else raw_scope

                codes = [c.strip().strip("'\"") for c in str(pt_codes_raw).split(",") if c.strip()]

                for pt_code in codes:
                    key = (smq_name, scope, pt_code)
                    if key in seen_keys:
                        continue
                    seen_keys.add(key)

                    pt_name = code_to_name.get(pt_code, pt_code)
                    category = code_to_category.get(pt_code)

                    batch.append(SMQTerm(
                        smq_name=smq_name,
                        smq_code=smq_code,
                        scope=scope,
                        pt_name=pt_name,
                        pt_name_upper=pt_name.upper(),
                        pt_code=pt_code,
                        category=category,
                        is_active=True,
                        reference_version=reference_version
                    ))

                    if len(batch) >= batch_size:
                        self.db.bulk_save_objects(batch)
                        self.db.commit()
                        total_inserted += len(batch)
                        batch = []
        else:
            # Fallback: Parse hierarchical outline from smq_pt_sheet directly
            stack = []  # (col_idx, smq_name)
            for row in ws_smq.iter_rows(values_only=True):
                if not row:
                    continue

                # Check for SMQ header
                smq_col = -1
                smq_name = None
                for c_idx, val in enumerate(row):
                    if val and "(smq)" in str(val).lower():
                        smq_col = c_idx
                        smq_name = str(val).strip()
                        break

                if smq_name:
                    while stack and stack[-1][0] >= smq_col:
                        stack.pop()
                    stack.append((smq_col, smq_name))
                    continue

                # Check for Scope (Narrow or Broad)
                scope_col = -1
                scope_val = None
                for c_idx, val in enumerate(row):
                    if val in ("Narrow", "Broad"):
                        scope_col = c_idx
                        scope_val = val
                        break

                if scope_val and scope_col >= 0 and len(row) > scope_col + 3:
                    pt_name = row[scope_col + 2]
                    pt_code = row[scope_col + 3]

                    if pt_name and pt_code:
                        pt_clean = str(pt_name).strip()
                        code_clean = str(pt_code).strip()
                        pt_upper = pt_clean.upper()

                        for _, s_name in stack:
                            # Add to Broad
                            b_key = (s_name, "Broad", code_clean)
                            if b_key not in seen_keys:
                                seen_keys.add(b_key)
                                batch.append(SMQTerm(
                                    smq_name=s_name,
                                    smq_code="",
                                    scope="Broad",
                                    pt_name=pt_clean,
                                    pt_name_upper=pt_upper,
                                    pt_code=code_clean,
                                    is_active=True,
                                    reference_version=reference_version
                                ))

                            # If narrow, add to Narrow
                            if scope_val == "Narrow":
                                n_key = (s_name, "Narrow", code_clean)
                                if n_key not in seen_keys:
                                    seen_keys.add(n_key)
                                    batch.append(SMQTerm(
                                        smq_name=s_name,
                                        smq_code="",
                                        scope="Narrow",
                                        pt_name=pt_clean,
                                        pt_name_upper=pt_upper,
                                        pt_code=code_clean,
                                        is_active=True,
                                        reference_version=reference_version
                                    ))

                if len(batch) >= batch_size:
                    self.db.bulk_save_objects(batch)
                    self.db.commit()
                    total_inserted += len(batch)
                    batch = []

        if batch:
            self.db.bulk_save_objects(batch)
            self.db.commit()
            total_inserted += len(batch)

        # Track dataset identity
        ds_id = f"ds_smq_{reference_version}_{file_hash[:8]}"
        dataset = self.db.query(Dataset).filter_by(id=ds_id).first()
        if not dataset:
            dataset = Dataset(
                id=ds_id,
                dataset_type="SMQ_REFERENCE",
                product_name="MedDRA Reference",
                reporting_period=f"MedDRA v{reference_version}",
                source_filename=file_name,
                source_file_hash=file_hash,
                total_events=total_inserted,
                status="ACTIVE"
            )
            self.db.add(dataset)
        else:
            dataset.total_events = total_inserted
            dataset.source_filename = file_name
            dataset.source_file_hash = file_hash

        audit = AuditLog(
            user_id="system_smq_loader",
            dataset_id=ds_id,
            action=f"Loaded authoritative MedDRA SMQ terminology v{reference_version} from {file_name}: {total_inserted} active terms."
        )
        self.db.add(audit)
        self.db.commit()

        return {
            "total_smq_terms": total_inserted,
            "reference_version": reference_version,
            "source_file": file_name,
            "file_hash": file_hash
        }
